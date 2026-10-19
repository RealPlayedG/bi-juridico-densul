"""
Backups dos dados, guardados no volume do Railway (PASTA_DADOS/backups).

Três momentos:
- diário: uma cópia por dia, mantida por DIAS_BACKUP dias (padrão 60);
- antes de cada atualização do sistema: ao subir uma versão nova do código, antes de
  qualquer mudança no banco, guarda uma cópia marcada com a versão anterior;
- antes de importar planilha e quando o administrador clica em "Fazer backup agora".

Cada backup é um .zip com:
- dados.json: todas as tabelas, completas (é o que o restaurar.py usa para voltar);
- planilha.xlsx: as mesmas informações em abas, para abrir no Excel / Google Sheets.
"""

import io
import json
import os
import threading
import time
import zipfile
from datetime import date, datetime

from sqlalchemy import inspect, select

PASTA = os.path.join(os.path.abspath(os.environ.get("PASTA_DADOS", "data")), "backups")
DIAS = int(os.environ.get("DIAS_BACKUP", "60"))
GUARDAR_OUTROS = 40  # backups de atualização, importação e manuais mantidos
ARQ_VERSAO = os.path.join(PASTA, "ultima_versao.txt")


def _tabelas():
    from db import meta
    return meta.sorted_tables


def _valor(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if hasattr(v, "quantize"):
        return float(v)
    return v


def _dump():
    from db import engine
    existentes = set(inspect(engine).get_table_names())
    dados = {}
    with engine.connect() as c:
        for t in _tabelas():
            if t.name in existentes:
                dados[t.name] = [{k: _valor(v) for k, v in r._mapping.items()} for r in c.execute(select(t))]
    return dados


def _xlsx(dados):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    nomes = {"processos": "Processos", "andamentos": "Andamentos", "custas": "Custas", "anexos": "Anexos",
             "historico": "Histórico", "usuarios": "Usuários", "lojas": "Lojas"}
    for t, linhas in dados.items():
        ws = wb.create_sheet(nomes.get(t, t)[:31])
        if not linhas:
            continue
        cols = [k for k in linhas[0] if k != "hash"]  # senha nunca vai para a planilha
        ws.append(cols)
        for cel in ws[1]:
            cel.font = Font(bold=True)
        for r in linhas:
            ws.append([r.get(k) for k in cols])
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def fazer(motivo="diario"):
    """Gera o .zip e devolve o caminho. motivo: diario, manual, antes-da-atualizacao..., antes-da-importacao."""
    from db import agora
    os.makedirs(PASTA, exist_ok=True)
    marca = agora()
    if motivo == "diario":
        nome = f"bi-juridico-{marca:%Y-%m-%d}-diario.zip"
    else:
        nome = f"bi-juridico-{marca:%Y-%m-%d_%H%M%S}-{motivo}.zip"
    destino = os.path.join(PASTA, nome)
    dados = _dump()
    tmp = destino + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("dados.json", json.dumps({"gerado_em": marca.isoformat(), "motivo": motivo, "tabelas": dados}, ensure_ascii=False))
        z.writestr("planilha.xlsx", _xlsx(dados))
    os.replace(tmp, destino)
    _limpar_antigos()
    return destino


def _limpar_antigos():
    arqs = sorted(f for f in os.listdir(PASTA) if f.endswith(".zip"))
    diarios = [f for f in arqs if f.endswith("-diario.zip")]
    outros = [f for f in arqs if not f.endswith("-diario.zip")]
    for velho in diarios[:-DIAS] + outros[:-GUARDAR_OUTROS]:
        try:
            os.remove(os.path.join(PASTA, velho))
        except OSError:
            pass


def listar():
    os.makedirs(PASTA, exist_ok=True)
    itens = sorted((f for f in os.listdir(PASTA) if f.endswith(".zip")), reverse=True)
    return [{"nome": f, "tamanho": os.path.getsize(os.path.join(PASTA, f))} for f in itens]


def antes_de_atualizar():
    """Na subida: se a versão do código mudou (RAILWAY_GIT_COMMIT_SHA), guarda os dados antes de qualquer migração."""
    versao = os.environ.get("RAILWAY_GIT_COMMIT_SHA", "")[:7]
    if not versao:
        return
    try:
        os.makedirs(PASTA, exist_ok=True)
        anterior = open(ARQ_VERSAO).read().strip() if os.path.exists(ARQ_VERSAO) else ""
        if anterior == versao:
            return
        from db import engine
        if anterior and inspect(engine).get_table_names():
            fazer(f"antes-da-atualizacao-{anterior}-para-{versao}")
        with open(ARQ_VERSAO, "w") as f:
            f.write(versao)
    except Exception as e:  # backup nunca derruba o sistema
        print("Falha no backup antes da atualização:", e, flush=True)


def rotina_diaria():
    if os.environ.get("BACKUP_AUTOMATICO", "1") != "1":
        return

    def loop():
        time.sleep(60)
        while True:
            try:
                from db import agora
                if not os.path.exists(os.path.join(PASTA, f"bi-juridico-{agora():%Y-%m-%d}-diario.zip")):
                    fazer("diario")
            except Exception as e:
                print("Falha no backup diário:", e, flush=True)
            time.sleep(3600)
    threading.Thread(target=loop, daemon=True, name="backup").start()
